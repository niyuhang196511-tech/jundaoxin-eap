"""文档解析（M12）：PDF/DOCX → 纯文本，供知识摄入。

可选依赖：pdf 解析 pypdf、docx 解析 python-docx（uv sync --extra docs 未装则报
ValueError 提示安装）。格式按扩展名分发；txt/md 直读。
"""

from __future__ import annotations


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名解析文件字节为纯文本。"""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf":
        return _pdf(data)
    if suffix in ("docx",):
        return _docx(data)
    if suffix in ("txt", "md", "markdown"):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"不支持的文档格式 .{suffix}（支持 pdf / docx / txt / md）")


def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("PDF 解析需要 pypdf：uv sync --extra docs") from e
    import io

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def _docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise ValueError("DOCX 解析需要 python-docx：uv sync --extra docs") from e
    import io

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:  # 表格逐行拼接
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
